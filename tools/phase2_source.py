#!/usr/bin/env python3
"""
Phase 2 — Listing Sourcing

Builds the regional XLSX of tourism business listings from DMO websites.

Usage:
    python tools/phase2_source.py --template
    python tools/phase2_source.py --region <key> --input <json_file>
    python tools/phase2_source.py --region <key> --input <json_file> --append

Process:
1. Run --template to get the JSON structure Claude should fill
2. Claude fetches DMO sites, extracts listings, cross-references trails
3. Claude saves the filled JSON
4. Run with --region and --input to build/append to the XLSX

Categories:
  food_drink, culture_heritage, attractions, wellness, outdoor, agriculture
"""

import argparse
import json
import os
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

REGIONS_FILE = os.path.join(os.path.dirname(__file__), "..", "data", "ny_regions.json")
OUTPUTS_DIR = os.path.join(os.path.dirname(__file__), "..", "outputs")

CATEGORIES = [
    "food_drink",
    "culture_heritage",
    "attractions",
    "wellness",
    "outdoor",
    "agriculture"
]

CATEGORY_LABELS = {
    "food_drink": "Food & Drink",
    "culture_heritage": "Culture & Heritage",
    "attractions": "Attractions",
    "wellness": "Wellness",
    "outdoor": "Outdoor",
    "agriculture": "Agriculture"
}

COLUMNS = [
    ("name", "Business Name", 30),
    ("address", "Address", 35),
    ("city", "City", 18),
    ("county", "County", 15),
    ("category", "Category", 18),
    ("subcategory", "Subcategory", 20),
    ("website", "Website", 35),
    ("phone", "Phone", 16),
    ("trail_member", "Trail Member", 14),
    ("trail_name", "Trail Name", 30),
    ("source_url", "Source URL", 35),
    ("notes", "Notes", 40),
]


def print_template():
    """Print the JSON template Claude should fill for each county."""
    template = {
        "county": "COUNTY_NAME",
        "source_dmo": "DMO_NAME",
        "source_url": "URL_FETCHED",
        "date_sourced": "YYYY-MM-DD",
        "listings": [
            {
                "name": "Business Name",
                "address": "123 Main St",
                "city": "City",
                "county": "COUNTY_NAME",
                "category": "food_drink | culture_heritage | attractions | wellness | outdoor | agriculture",
                "subcategory": "e.g., brewery, museum, hiking outfitter",
                "website": "https://...",
                "phone": "555-555-5555",
                "trail_member": False,
                "trail_name": None,
                "source_url": "URL where listing was found",
                "notes": "Any relevant notes"
            }
        ]
    }
    print(json.dumps(template, indent=2))
    print()
    print("INSTRUCTIONS:")
    print("1. Fill one JSON file per county")
    print("2. category must be one of:", ", ".join(CATEGORIES))
    print("3. Set trail_member: true only for confirmed trail stops")
    print("4. If trail_member is true, include trail_name")
    print("5. An operator on multiple trails gets one row per trail")


def load_regions():
    with open(REGIONS_FILE, "r") as f:
        return json.load(f)


def validate_listings(listings_data):
    """Validate the listings JSON structure."""
    errors = []

    if "county" not in listings_data:
        errors.append("Missing 'county' field")
    if "listings" not in listings_data:
        errors.append("Missing 'listings' array")
        return errors

    for i, listing in enumerate(listings_data["listings"]):
        prefix = f"Listing {i + 1}"
        if not listing.get("name"):
            errors.append(f"{prefix}: missing 'name'")
        if not listing.get("city"):
            errors.append(f"{prefix} ({listing.get('name', '?')}): missing 'city'")
        if listing.get("category") and listing["category"] not in CATEGORIES:
            errors.append(f"{prefix} ({listing.get('name', '?')}): invalid category '{listing['category']}'. Must be one of: {', '.join(CATEGORIES)}")
        if listing.get("trail_member") and not listing.get("trail_name"):
            errors.append(f"{prefix} ({listing.get('name', '?')}): trail_member is true but trail_name is missing")

    return errors


def create_workbook(region_name):
    """Create a new styled workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Listings"

    # Header styling
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Write headers
    for col_idx, (key, label, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = width

    # Freeze top row
    ws.freeze_panes = "A2"

    # Add metadata sheet
    meta_ws = wb.create_sheet("Metadata")
    meta_ws["A1"] = "Region"
    meta_ws["B1"] = region_name
    meta_ws["A2"] = "Generated"
    meta_ws["B2"] = "Upstate Regional Guide System"
    meta_ws["A3"] = "Categories"
    for i, (cat, label) in enumerate(CATEGORY_LABELS.items()):
        meta_ws[f"A{4 + i}"] = cat
        meta_ws[f"B{4 + i}"] = label

    return wb


def add_listings_to_workbook(wb, listings_data):
    """Add listings from JSON data to the workbook."""
    ws = wb["Listings"]
    start_row = ws.max_row + 1

    # Category colors for visual distinction
    category_colors = {
        "food_drink": "E8F5E9",
        "culture_heritage": "FFF3E0",
        "attractions": "E3F2FD",
        "wellness": "F3E5F5",
        "outdoor": "E0F2F1",
        "agriculture": "FFF9C4"
    }

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    trail_font = Font(name="Calibri", bold=True, color="1B5E20")
    count = 0

    for listing in listings_data.get("listings", []):
        row = start_row + count
        for col_idx, (key, label, width) in enumerate(COLUMNS, 1):
            value = listing.get(key, "")
            if isinstance(value, bool):
                value = "Yes" if value else "No"
            elif value is None:
                value = ""

            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Color by category
            cat = listing.get("category", "")
            if cat in category_colors:
                cell.fill = PatternFill(start_color=category_colors[cat], end_color=category_colors[cat], fill_type="solid")

            # Bold trail members
            if listing.get("trail_member"):
                cell.font = trail_font

        count += 1

    return count


def build_xlsx(region_key, input_file, append=False):
    """Build or append to the regional XLSX."""
    data = load_regions()
    if region_key not in data["regions"]:
        print(f"Error: Region '{region_key}' not found.", file=sys.stderr)
        sys.exit(1)

    region = data["regions"][region_key]
    os.makedirs(OUTPUTS_DIR, exist_ok=True)
    output_path = os.path.join(OUTPUTS_DIR, f"{region_key}-listings.xlsx")

    # Load input JSON
    with open(input_file, "r") as f:
        listings_data = json.load(f)

    # Validate
    errors = validate_listings(listings_data)
    if errors:
        print("Validation errors:", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        sys.exit(1)

    # Create or load workbook
    if append and os.path.exists(output_path):
        wb = load_workbook(output_path)
        print(f"Appending to existing file: {output_path}")
    else:
        wb = create_workbook(region["name"])
        if append:
            print(f"Note: No existing file found, creating new: {output_path}")

    # Add listings
    count = add_listings_to_workbook(wb, listings_data)

    # Update metadata
    meta_ws = wb["Metadata"]
    # Find next empty row in metadata for source tracking
    meta_row = 1
    while meta_ws.cell(row=meta_row, column=4).value:
        meta_row += 1
    if meta_row == 1:
        meta_ws["D1"] = "Source County"
        meta_ws["E1"] = "Source DMO"
        meta_ws["F1"] = "Listings Count"
        meta_ws["G1"] = "Date Sourced"
        meta_row = 2

    meta_ws.cell(row=meta_row, column=4, value=listings_data.get("county", ""))
    meta_ws.cell(row=meta_row, column=5, value=listings_data.get("source_dmo", ""))
    meta_ws.cell(row=meta_row, column=6, value=count)
    meta_ws.cell(row=meta_row, column=7, value=listings_data.get("date_sourced", ""))

    # Save
    wb.save(output_path)

    # Summary
    total_rows = wb["Listings"].max_row - 1  # minus header
    print(f"\nListings added: {count}")
    print(f"Total listings in file: {total_rows}")
    print(f"County: {listings_data.get('county', 'Unknown')}")
    print(f"Source: {listings_data.get('source_dmo', 'Unknown')}")
    print(f"Output: {output_path}")

    # Category breakdown
    categories = {}
    trail_members = 0
    for listing in listings_data.get("listings", []):
        cat = listing.get("category", "uncategorized")
        categories[cat] = categories.get(cat, 0) + 1
        if listing.get("trail_member"):
            trail_members += 1

    print(f"\nCategory breakdown (this batch):")
    for cat, cnt in sorted(categories.items()):
        label = CATEGORY_LABELS.get(cat, cat)
        print(f"  {label}: {cnt}")
    print(f"  Trail members: {trail_members}")


def main():
    parser = argparse.ArgumentParser(description="Phase 2: Listing Sourcing")
    parser.add_argument("--region", type=str, help="Region key")
    parser.add_argument("--input", type=str, help="Input JSON file with listings")
    parser.add_argument("--append", action="store_true", help="Append to existing XLSX")
    parser.add_argument("--template", action="store_true", help="Print JSON template")
    args = parser.parse_args()

    if args.template:
        print_template()
        return

    if not args.region or not args.input:
        parser.print_help()
        print("\nUsage examples:")
        print("  python tools/phase2_source.py --template")
        print("  python tools/phase2_source.py --region finger_lakes --input /tmp/schuyler.json")
        print("  python tools/phase2_source.py --region finger_lakes --input /tmp/ontario.json --append")
        sys.exit(1)

    if not os.path.exists(args.input):
        print(f"Error: Input file not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    build_xlsx(args.region, args.input, args.append)


if __name__ == "__main__":
    main()
